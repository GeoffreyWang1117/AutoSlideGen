"""
Template management system.
Allows users to use custom PowerPoint templates.
"""

import logging
from pathlib import Path
from typing import Optional, Dict, Any
import shutil

from pptx import Presentation

logger = logging.getLogger(__name__)


class TemplateManager:
    """Manages PowerPoint templates."""

    def __init__(self, templates_dir: Optional[str] = None):
        """
        Initialize template manager.

        Args:
            templates_dir: Directory containing template files
        """
        if templates_dir is None:
            templates_dir = "./templates"

        self.templates_dir = Path(templates_dir)
        self.templates_dir.mkdir(parents=True, exist_ok=True)
        self.logger = logging.getLogger(self.__class__.__name__)

        self.templates = self._discover_templates()

    def _discover_templates(self) -> Dict[str, Path]:
        """
        Discover available templates.

        Returns:
            Dictionary mapping template names to paths
        """
        templates = {}

        for template_file in self.templates_dir.glob("*.pptx"):
            template_name = template_file.stem
            templates[template_name] = template_file
            self.logger.debug(f"Discovered template: {template_name}")

        if templates:
            self.logger.info(f"Found {len(templates)} templates")

        return templates

    def list_templates(self) -> list:
        """
        List available template names.

        Returns:
            List of template names
        """
        return list(self.templates.keys())

    def get_template_path(self, template_name: str) -> Optional[Path]:
        """
        Get path to template file.

        Args:
            template_name: Name of template

        Returns:
            Path to template or None if not found
        """
        return self.templates.get(template_name)

    def load_template(self, template_name: str) -> Optional[Presentation]:
        """
        Load template presentation.

        Args:
            template_name: Name of template to load

        Returns:
            Presentation object or None if not found
        """
        template_path = self.get_template_path(template_name)

        if template_path is None:
            self.logger.error(f"Template not found: {template_name}")
            return None

        try:
            prs = Presentation(str(template_path))
            self.logger.info(f"Loaded template: {template_name}")
            return prs

        except Exception as e:
            self.logger.error(f"Failed to load template {template_name}: {e}")
            return None

    def add_template(
        self,
        template_path: str,
        template_name: Optional[str] = None
    ) -> bool:
        """
        Add a new template.

        Args:
            template_path: Path to PPTX template file
            template_name: Name for the template (uses filename if None)

        Returns:
            True if successful, False otherwise
        """
        source_path = Path(template_path)

        if not source_path.exists():
            self.logger.error(f"Template file not found: {template_path}")
            return False

        if template_name is None:
            template_name = source_path.stem

        dest_path = self.templates_dir / f"{template_name}.pptx"

        try:
            shutil.copy2(source_path, dest_path)
            self.templates[template_name] = dest_path
            self.logger.info(f"Added template: {template_name}")
            return True

        except Exception as e:
            self.logger.error(f"Failed to add template: {e}")
            return False

    def remove_template(self, template_name: str) -> bool:
        """
        Remove a template.

        Args:
            template_name: Name of template to remove

        Returns:
            True if successful, False otherwise
        """
        if template_name not in self.templates:
            self.logger.error(f"Template not found: {template_name}")
            return False

        template_path = self.templates[template_name]

        try:
            template_path.unlink()
            del self.templates[template_name]
            self.logger.info(f"Removed template: {template_name}")
            return True

        except Exception as e:
            self.logger.error(f"Failed to remove template: {e}")
            return False

    def get_template_info(self, template_name: str) -> Optional[Dict[str, Any]]:
        """
        Get information about a template.

        Args:
            template_name: Name of template

        Returns:
            Dictionary with template information or None
        """
        template_path = self.get_template_path(template_name)

        if template_path is None:
            return None

        try:
            prs = Presentation(str(template_path))

            info = {
                'name': template_name,
                'path': str(template_path),
                'slide_width': prs.slide_width,
                'slide_height': prs.slide_height,
                'num_layouts': len(prs.slide_layouts),
                'layouts': []
            }

            for i, layout in enumerate(prs.slide_layouts):
                info['layouts'].append({
                    'index': i,
                    'name': layout.name,
                    'num_placeholders': len(layout.placeholders)
                })

            return info

        except Exception as e:
            self.logger.error(f"Failed to get template info: {e}")
            return None


class CustomizableTemplateBuilder:
    """PPT builder that uses custom templates."""

    def __init__(
        self,
        template_manager: TemplateManager,
        template_name: Optional[str] = None
    ):
        """
        Initialize customizable builder.

        Args:
            template_manager: TemplateManager instance
            template_name: Name of template to use (None for blank)
        """
        self.template_manager = template_manager
        self.template_name = template_name
        self.logger = logging.getLogger(self.__class__.__name__)

        # Layout mapping (can be customized)
        self.layout_mapping = {
            'title': 0,      # Title slide layout
            'content': 1,    # Title and content layout
            'section': 2,    # Section header layout
            'blank': 6       # Blank layout
        }

    def create_presentation(self) -> Presentation:
        """
        Create presentation from template or blank.

        Returns:
            Presentation object
        """
        if self.template_name:
            prs = self.template_manager.load_template(self.template_name)

            if prs is None:
                self.logger.warning(
                    f"Template '{self.template_name}' not found, using blank presentation"
                )
                prs = Presentation()
        else:
            prs = Presentation()

        return prs

    def set_layout_mapping(self, mapping: Dict[str, int]):
        """
        Set custom layout index mapping.

        Args:
            mapping: Dictionary mapping slide types to layout indices
        """
        self.layout_mapping.update(mapping)
        self.logger.info("Updated layout mapping")

    def get_layout_index(self, slide_type: str) -> int:
        """
        Get layout index for slide type.

        Args:
            slide_type: Type of slide

        Returns:
            Layout index
        """
        return self.layout_mapping.get(slide_type, 1)  # Default to content layout
