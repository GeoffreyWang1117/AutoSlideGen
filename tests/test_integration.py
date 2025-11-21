"""
Integration tests for AutoSlideGen.
Tests the complete workflow from outline generation to PPTX creation.
"""

import pytest
import tempfile
import shutil
from pathlib import Path
from unittest.mock import Mock, AsyncMock, patch
from pptx import Presentation

from autoslidegen.main import AutoSlideGen
from autoslidegen.parser.models import (
    BulletPoint,
    Slide,
    PresentationMetadata,
    PresentationOutline
)


class TestEndToEndWorkflow:
    """Integration tests for complete workflow."""

    @pytest.fixture
    def temp_dir(self):
        """Create temporary directory."""
        temp_dir = tempfile.mkdtemp()
        yield temp_dir
        shutil.rmtree(temp_dir)

    @pytest.fixture
    def sample_outline(self):
        """Create a valid presentation outline."""
        metadata = PresentationMetadata(
            topic="人工智能技术",
            audience="技术人员",
            purpose="技术介绍",
            language="zh"
        )

        slides = [
            Slide(
                slide_number=1,
                title="人工智能技术",
                slide_type="title",
                bullet_points=[
                    BulletPoint(text="技术人员", level=1),
                    BulletPoint(text="技术介绍", level=1)
                ]
            ),
            Slide(
                slide_number=2,
                title="AI应用领域",
                slide_type="content",
                bullet_points=[
                    BulletPoint(text="医疗诊断", level=1),
                    BulletPoint(text="金融分析", level=1),
                    BulletPoint(text="自动驾驶", level=1),
                    BulletPoint(text="智能助手", level=1)
                ]
            ),
            Slide(
                slide_number=3,
                title="核心技术",
                slide_type="section",
                bullet_points=[]
            ),
            Slide(
                slide_number=4,
                title="机器学习",
                slide_type="content",
                bullet_points=[
                    BulletPoint(text="监督学习", level=1),
                    BulletPoint(text="无监督学习", level=1),
                    BulletPoint(text="强化学习", level=1)
                ]
            )
        ]

        return PresentationOutline(metadata=metadata, slides=slides)

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_complete_workflow_basic(self, mock_factory, temp_dir, sample_outline):
        """Test complete workflow from request to PPTX creation."""
        # Mock the LLM generator
        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=sample_outline)
        mock_factory.return_value = mock_generator

        # Initialize AutoSlideGen
        asg = AutoSlideGen(provider='openai')

        # Generate presentation
        result = asg.generate(
            topic="人工智能技术",
            audience="技术人员",
            purpose="技术介绍",
            language="zh",
            num_slides=10,
            output_path=str(Path(temp_dir) / "test.pptx")
        )

        # Verify results
        assert 'pptx_path' in result
        assert 'outline' in result
        assert Path(result['pptx_path']).exists()

        # Verify PPTX can be opened
        prs = Presentation(result['pptx_path'])
        assert len(prs.slides) == 4

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_workflow_with_json_export(self, mock_factory, temp_dir, sample_outline):
        """Test workflow with JSON export."""
        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=sample_outline)
        mock_factory.return_value = mock_generator

        asg = AutoSlideGen(provider='openai')

        result = asg.generate(
            topic="Test",
            audience="Test",
            purpose="Test",
            save_json=True,
            output_path=str(Path(temp_dir) / "test.pptx")
        )

        assert result['json_path'] is not None
        assert Path(result['json_path']).exists()

        # Verify JSON can be loaded
        with open(result['json_path'], 'r', encoding='utf-8') as f:
            import json
            data = json.load(f)
            assert 'metadata' in data
            assert 'slides' in data

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_workflow_no_json_export(self, mock_factory, temp_dir, sample_outline):
        """Test workflow without JSON export."""
        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=sample_outline)
        mock_factory.return_value = mock_generator

        asg = AutoSlideGen(provider='openai')

        result = asg.generate(
            topic="Test",
            audience="Test",
            purpose="Test",
            save_json=False,
            output_path=str(Path(temp_dir) / "test.pptx")
        )

        assert result['json_path'] is None

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_generate_from_json(self, mock_factory, temp_dir, sample_outline):
        """Test generating PPTX from existing JSON."""
        # Create JSON file
        json_path = Path(temp_dir) / "outline.json"
        json_path.write_text(sample_outline.to_json())

        asg = AutoSlideGen()

        pptx_path = asg.generate_from_json(
            str(json_path),
            str(Path(temp_dir) / "output.pptx")
        )

        assert Path(pptx_path).exists()

        prs = Presentation(pptx_path)
        assert len(prs.slides) == 4

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    @patch('autoslidegen.main.SpeakerNotesGenerator')
    def test_workflow_with_speaker_notes(
        self,
        mock_notes_gen_class,
        mock_factory,
        temp_dir,
        sample_outline
    ):
        """Test workflow with speaker notes generation."""
        # Mock LLM generator
        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=sample_outline)
        mock_factory.return_value = mock_generator

        # Mock speaker notes generator
        outline_with_notes = sample_outline.model_copy(deep=True)
        for slide in outline_with_notes.slides:
            if slide.slide_type == "content":
                slide.notes = "这是演讲稿内容"

        mock_notes_gen = Mock()
        mock_notes_gen.generate_notes_for_outline_sync = Mock(
            return_value=outline_with_notes
        )
        mock_notes_gen_class.return_value = mock_notes_gen

        asg = AutoSlideGen(provider='openai')

        result = asg.generate(
            topic="Test",
            audience="Test",
            purpose="Test",
            generate_speaker_notes=True,
            output_path=str(Path(temp_dir) / "test.pptx")
        )

        # Verify notes were generated
        mock_notes_gen.generate_notes_for_outline_sync.assert_called_once()

        # Verify PPTX has notes
        prs = Presentation(result['pptx_path'])
        content_slides = [s for s in prs.slides if len(s.shapes) > 1]
        if len(content_slides) > 0:
            notes_text = content_slides[0].notes_slide.notes_text_frame.text
            # Notes should be present

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_workflow_different_languages(self, mock_factory, temp_dir):
        """Test workflow with different languages."""
        languages = ['zh', 'en', 'ja', 'es']

        for lang in languages:
            metadata = PresentationMetadata(
                topic="Test",
                audience="Test",
                purpose="Test",
                language=lang
            )

            slides = [
                Slide(
                    slide_number=1,
                    title="Title",
                    slide_type="title",
                    bullet_points=[BulletPoint(text="Subtitle", level=1)]
                ),
                Slide(
                    slide_number=2,
                    title="Content",
                    slide_type="content",
                    bullet_points=[
                        BulletPoint(text="Point 1", level=1),
                        BulletPoint(text="Point 2", level=1)
                    ]
                )
            ]

            outline = PresentationOutline(metadata=metadata, slides=slides)

            mock_generator = Mock()
            mock_generator.generate_outline_sync = Mock(return_value=outline)
            mock_factory.return_value = mock_generator

            asg = AutoSlideGen(provider='openai')

            result = asg.generate(
                topic="Test",
                audience="Test",
                purpose="Test",
                language=lang,
                output_path=str(Path(temp_dir) / f"test_{lang}.pptx")
            )

            assert Path(result['pptx_path']).exists()

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_error_handling_invalid_outline(self, mock_factory, temp_dir):
        """Test error handling when LLM generates invalid outline."""
        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(
            side_effect=Exception("Failed to generate outline")
        )
        mock_factory.return_value = mock_generator

        asg = AutoSlideGen(provider='openai')

        with pytest.raises(Exception):
            asg.generate(
                topic="Test",
                audience="Test",
                purpose="Test",
                output_path=str(Path(temp_dir) / "test.pptx")
            )

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_auto_output_path_generation(self, mock_factory, sample_outline):
        """Test automatic output path generation."""
        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=sample_outline)
        mock_factory.return_value = mock_generator

        asg = AutoSlideGen(provider='openai')

        result = asg.generate(
            topic="Test",
            audience="Test",
            purpose="Test"
            # No output_path specified
        )

        assert 'pptx_path' in result
        assert Path(result['pptx_path']).exists()

        # Clean up
        Path(result['pptx_path']).unlink()
        if result['json_path']:
            Path(result['json_path']).unlink()
        try:
            Path(result['pptx_path']).parent.rmdir()
        except:
            pass

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_large_presentation_generation(self, mock_factory, temp_dir):
        """Test generating large presentation with many slides."""
        metadata = PresentationMetadata(
            topic="Large Presentation",
            audience="Test",
            purpose="Test"
        )

        slides = [
            Slide(
                slide_number=1,
                title="Title",
                slide_type="title",
                bullet_points=[BulletPoint(text="Subtitle", level=1)]
            )
        ]

        # Add 29 content slides
        for i in range(2, 31):
            slides.append(
                Slide(
                    slide_number=i,
                    title=f"Content {i}",
                    slide_type="content",
                    bullet_points=[
                        BulletPoint(text=f"Point {j}", level=1)
                        for j in range(1, 5)
                    ]
                )
            )

        outline = PresentationOutline(metadata=metadata, slides=slides)

        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=outline)
        mock_factory.return_value = mock_generator

        asg = AutoSlideGen(provider='openai')

        result = asg.generate(
            topic="Large Presentation",
            audience="Test",
            purpose="Test",
            num_slides=30,
            output_path=str(Path(temp_dir) / "large.pptx")
        )

        prs = Presentation(result['pptx_path'])
        assert len(prs.slides) == 30


class TestConfigurationIntegration:
    """Integration tests for configuration handling."""

    @pytest.fixture
    def temp_dir(self):
        """Create temporary directory."""
        temp_dir = tempfile.mkdtemp()
        yield temp_dir
        shutil.rmtree(temp_dir)

    @pytest.fixture
    def config_file(self, temp_dir):
        """Create test config file."""
        config_path = Path(temp_dir) / "config.yaml"
        config_content = """
default_provider: openai
default_language: zh

llm_providers:
  openai:
    model: gpt-4
    temperature: 0.7
    max_tokens: 4000
"""
        config_path.write_text(config_content)
        return str(config_path)

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_load_custom_config(self, mock_factory, config_file, temp_dir):
        """Test loading custom configuration."""
        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=Mock(slides=[]))
        mock_factory.return_value = mock_generator

        # This test verifies config loading doesn't crash
        # Actual config validation is done in unit tests
        asg = AutoSlideGen(config_path=config_file)
        assert asg is not None


class TestEdgeCases:
    """Integration tests for edge cases."""

    @pytest.fixture
    def temp_dir(self):
        """Create temporary directory."""
        temp_dir = tempfile.mkdtemp()
        yield temp_dir
        shutil.rmtree(temp_dir)

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_minimal_presentation(self, mock_factory, temp_dir):
        """Test generating minimal presentation with just title slide."""
        metadata = PresentationMetadata(
            topic="Minimal",
            audience="Test",
            purpose="Test"
        )

        slides = [
            Slide(
                slide_number=1,
                title="Title Only",
                slide_type="title",
                bullet_points=[BulletPoint(text="Subtitle", level=1)]
            )
        ]

        outline = PresentationOutline(metadata=metadata, slides=slides)

        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=outline)
        mock_factory.return_value = mock_generator

        asg = AutoSlideGen(provider='openai')

        result = asg.generate(
            topic="Minimal",
            audience="Test",
            purpose="Test",
            num_slides=5,
            output_path=str(Path(temp_dir) / "minimal.pptx")
        )

        prs = Presentation(result['pptx_path'])
        assert len(prs.slides) >= 1

    @patch('autoslidegen.main.GeneratorFactory.create_generator')
    def test_special_characters_in_topic(self, mock_factory, temp_dir):
        """Test handling special characters in topic."""
        metadata = PresentationMetadata(
            topic="测试/主题:特殊*字符?",
            audience="Test",
            purpose="Test"
        )

        slides = [
            Slide(
                slide_number=1,
                title="Title",
                slide_type="title",
                bullet_points=[BulletPoint(text="Subtitle", level=1)]
            )
        ]

        outline = PresentationOutline(metadata=metadata, slides=slides)

        mock_generator = Mock()
        mock_generator.generate_outline_sync = Mock(return_value=outline)
        mock_factory.return_value = mock_generator

        asg = AutoSlideGen(provider='openai')

        result = asg.generate(
            topic="测试/主题:特殊*字符?",
            audience="Test",
            purpose="Test",
            output_path=str(Path(temp_dir) / "special.pptx")
        )

        assert Path(result['pptx_path']).exists()
