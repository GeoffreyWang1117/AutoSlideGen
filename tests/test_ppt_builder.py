"""
Tests for PPT builder module.
"""

import pytest
import tempfile
import shutil
from pathlib import Path
from unittest.mock import Mock, patch

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from autoslidegen.ppt_builder.builder import PPTBuilder
from autoslidegen.parser.models import (
    BulletPoint,
    Slide,
    PresentationMetadata,
    PresentationOutline
)


class TestPPTBuilder:
    """Tests for PPTBuilder."""

    @pytest.fixture
    def builder(self):
        """Create a basic PPT builder instance."""
        config = {
            'fonts': {
                'title': {'name': 'Arial', 'size': 32, 'bold': True, 'color': '1F4E78'},
                'subtitle': {'name': 'Arial', 'size': 18, 'bold': False, 'color': '5A5A5A'},
                'bullet': {'name': 'Arial', 'size': 16, 'bold': False, 'color': '333333'}
            }
        }
        return PPTBuilder(config)

    @pytest.fixture
    def simple_outline(self):
        """Create a simple presentation outline."""
        metadata = PresentationMetadata(
            topic="测试演示",
            audience="测试观众",
            purpose="测试目的",
            language="zh"
        )

        slides = [
            Slide(
                slide_number=1,
                title="标题页",
                slide_type="title",
                bullet_points=[BulletPoint(text="副标题", level=1)]
            ),
            Slide(
                slide_number=2,
                title="内容页",
                slide_type="content",
                bullet_points=[
                    BulletPoint(text="要点一", level=1),
                    BulletPoint(text="要点二", level=1),
                    BulletPoint(text="详细说明", level=2)
                ]
            ),
            Slide(
                slide_number=3,
                title="章节页",
                slide_type="section",
                bullet_points=[]
            )
        ]

        return PresentationOutline(metadata=metadata, slides=slides)

    @pytest.fixture
    def temp_output_dir(self):
        """Create a temporary output directory."""
        temp_dir = tempfile.mkdtemp()
        yield temp_dir
        shutil.rmtree(temp_dir)

    def test_builder_initialization(self):
        """Test PPT builder initialization."""
        builder = PPTBuilder()
        assert builder.config == {}
        assert builder.image_map == {}

        config = {'fonts': {'title': {'size': 32}}}
        builder = PPTBuilder(config)
        assert builder.config == config
        assert builder.fonts == {'title': {'size': 32}}

    def test_hex_to_rgb_conversion(self, builder):
        """Test hex color to RGB conversion."""
        # Test with # prefix
        rgb = builder._hex_to_rgb("#1F4E78")
        assert isinstance(rgb, RGBColor)
        assert rgb.rgb == (31, 78, 120)

        # Test without # prefix
        rgb = builder._hex_to_rgb("FF0000")
        assert rgb.rgb == (255, 0, 0)

        # Test black
        rgb = builder._hex_to_rgb("000000")
        assert rgb.rgb == (0, 0, 0)

        # Test white
        rgb = builder._hex_to_rgb("FFFFFF")
        assert rgb.rgb == (255, 255, 255)

    def test_build_simple_presentation(self, builder, simple_outline, temp_output_dir):
        """Test building a simple presentation."""
        output_path = str(Path(temp_output_dir) / "test.pptx")
        result_path = builder.build(simple_outline, output_path)

        assert result_path == output_path
        assert Path(output_path).exists()

        # Verify the presentation can be opened
        prs = Presentation(output_path)
        assert len(prs.slides) == 3

    def test_build_with_auto_generated_filename(self, builder, simple_outline):
        """Test building with auto-generated filename."""
        result_path = builder.build(simple_outline)

        assert Path(result_path).exists()
        assert result_path.endswith('.pptx')
        assert '测试演示' in result_path or 'output' in result_path

        # Cleanup
        Path(result_path).unlink()
        # Try to remove output directory if empty
        try:
            Path(result_path).parent.rmdir()
        except:
            pass

    def test_build_creates_output_directory(self, builder, simple_outline, temp_output_dir):
        """Test that build creates output directory if it doesn't exist."""
        output_path = str(Path(temp_output_dir) / "subdir" / "test.pptx")
        result_path = builder.build(simple_outline, output_path)

        assert Path(output_path).exists()
        assert Path(output_path).parent.exists()

    def test_build_title_slide(self, builder, simple_outline, temp_output_dir):
        """Test building title slide."""
        output_path = str(Path(temp_output_dir) / "test.pptx")
        builder.build(simple_outline, output_path)

        prs = Presentation(output_path)
        title_slide = prs.slides[0]

        # Verify title
        assert title_slide.shapes.title.text == "标题页"

    def test_build_content_slide(self, builder, simple_outline, temp_output_dir):
        """Test building content slide with bullets."""
        output_path = str(Path(temp_output_dir) / "test.pptx")
        builder.build(simple_outline, output_path)

        prs = Presentation(output_path)
        content_slide = prs.slides[1]

        # Verify title
        assert content_slide.shapes.title.text == "内容页"

        # Verify bullet points exist
        # Note: Specific bullet verification depends on layout
        assert len(content_slide.shapes) >= 2

    def test_build_section_slide(self, builder, simple_outline, temp_output_dir):
        """Test building section slide."""
        output_path = str(Path(temp_output_dir) / "test.pptx")
        builder.build(simple_outline, output_path)

        prs = Presentation(output_path)
        section_slide = prs.slides[2]

        # Verify title
        assert section_slide.shapes.title.text == "章节页"

    def test_build_with_speaker_notes(self, builder, temp_output_dir):
        """Test building presentation with speaker notes."""
        metadata = PresentationMetadata(
            topic="测试",
            audience="观众",
            purpose="目的"
        )

        slides = [
            Slide(
                slide_number=1,
                title="标题",
                slide_type="title",
                bullet_points=[BulletPoint(text="副标题", level=1)]
            ),
            Slide(
                slide_number=2,
                title="内容",
                slide_type="content",
                bullet_points=[
                    BulletPoint(text="要点", level=1)
                ],
                notes="这是演讲稿内容"
            )
        ]

        outline = PresentationOutline(metadata=metadata, slides=slides)
        output_path = str(Path(temp_output_dir) / "test_notes.pptx")
        builder.build(outline, output_path)

        prs = Presentation(output_path)
        content_slide = prs.slides[1]

        # Verify notes
        notes_slide = content_slide.notes_slide
        notes_text = notes_slide.notes_text_frame.text
        assert "这是演讲稿内容" in notes_text

    def test_build_with_images(self, builder, simple_outline, temp_output_dir):
        """Test building presentation with images."""
        # Create a dummy image file
        import io
        from PIL import Image

        img = Image.new('RGB', (100, 100), color='red')
        image_path = str(Path(temp_output_dir) / "test_image.jpg")
        img.save(image_path)

        # Add image to image map
        image_map = {2: image_path}

        output_path = str(Path(temp_output_dir) / "test_images.pptx")
        builder.build(simple_outline, output_path, image_map=image_map)

        assert Path(output_path).exists()

        prs = Presentation(output_path)
        # Should have 3 slides
        assert len(prs.slides) == 3

    def test_build_with_custom_slide_size(self, simple_outline, temp_output_dir):
        """Test building with custom slide size."""
        config = {
            'template': {
                'slide_width': 13.333,
                'slide_height': 7.5
            }
        }

        builder = PPTBuilder(config)
        output_path = str(Path(temp_output_dir) / "test_custom_size.pptx")
        builder.build(simple_outline, output_path)

        prs = Presentation(output_path)
        assert prs.slide_width == Inches(13.333)
        assert prs.slide_height == Inches(7.5)

    def test_build_from_json(self, builder, simple_outline, temp_output_dir):
        """Test building from JSON file."""
        # Save outline to JSON
        json_path = str(Path(temp_output_dir) / "outline.json")
        with open(json_path, 'w', encoding='utf-8') as f:
            f.write(simple_outline.to_json())

        # Build from JSON
        output_path = str(Path(temp_output_dir) / "from_json.pptx")
        result_path = builder.build_from_json(json_path, output_path)

        assert result_path == output_path
        assert Path(output_path).exists()

        prs = Presentation(output_path)
        assert len(prs.slides) == 3

    def test_build_with_multilevel_bullets(self, builder, temp_output_dir):
        """Test building with multi-level bullet points."""
        metadata = PresentationMetadata(
            topic="测试",
            audience="观众",
            purpose="目的"
        )

        slides = [
            Slide(
                slide_number=1,
                title="标题",
                slide_type="title",
                bullet_points=[BulletPoint(text="副标题", level=1)]
            ),
            Slide(
                slide_number=2,
                title="多级要点",
                slide_type="content",
                bullet_points=[
                    BulletPoint(text="一级要点1", level=1),
                    BulletPoint(text="二级要点1", level=2),
                    BulletPoint(text="二级要点2", level=2),
                    BulletPoint(text="一级要点2", level=1),
                    BulletPoint(text="三级要点", level=3)
                ]
            )
        ]

        outline = PresentationOutline(metadata=metadata, slides=slides)
        output_path = str(Path(temp_output_dir) / "test_multilevel.pptx")
        builder.build(outline, output_path)

        assert Path(output_path).exists()

    def test_build_with_different_languages(self, builder, temp_output_dir):
        """Test building presentations in different languages."""
        languages = [
            ('zh', '中文演示', '观众'),
            ('en', 'English Presentation', 'Audience'),
            ('ja', '日本語プレゼンテーション', '聴衆'),
            ('es', 'Presentación en Español', 'Audiencia')
        ]

        for lang, topic, audience in languages:
            metadata = PresentationMetadata(
                topic=topic,
                audience=audience,
                purpose="Test",
                language=lang
            )

            slides = [
                Slide(
                    slide_number=1,
                    title=topic,
                    slide_type="title",
                    bullet_points=[BulletPoint(text="Subtitle", level=1)]
                ),
                Slide(
                    slide_number=2,
                    title="Content",
                    slide_type="content",
                    bullet_points=[
                        BulletPoint(text="Point 1", level=1),
                        BulletPoint(text="Point 2", level=1)
                    ]
                )
            ]

            outline = PresentationOutline(metadata=metadata, slides=slides)
            output_path = str(Path(temp_output_dir) / f"test_{lang}.pptx")
            builder.build(outline, output_path)

            assert Path(output_path).exists()

    def test_filename_sanitization(self, builder, temp_output_dir):
        """Test that special characters in filenames are sanitized."""
        metadata = PresentationMetadata(
            topic="测试/演示:特殊*字符?",
            audience="观众",
            purpose="目的"
        )

        slides = [
            Slide(
                slide_number=1,
                title="标题",
                slide_type="title",
                bullet_points=[BulletPoint(text="副标题", level=1)]
            )
        ]

        outline = PresentationOutline(metadata=metadata, slides=slides)

        # Let it auto-generate filename
        result_path = builder.build(outline)

        # Should create file without special characters
        assert Path(result_path).exists()

        # Cleanup
        Path(result_path).unlink()
        try:
            Path(result_path).parent.rmdir()
        except:
            pass

    def test_apply_font_style(self, builder):
        """Test font styling application."""
        from pptx import Presentation
        from pptx.enum.text import PP_ALIGN

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        title.text = "Test"

        font_config = {
            'name': 'Calibri',
            'size': 24,
            'bold': True,
            'color': 'FF0000'
        }

        builder._apply_font_style(title.text_frame, font_config, PP_ALIGN.CENTER)

        # Verify alignment
        for para in title.text_frame.paragraphs:
            assert para.alignment == PP_ALIGN.CENTER

    def test_add_image_to_slide_error_handling(self, builder, temp_output_dir):
        """Test error handling when adding invalid image."""
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])

        # Try to add non-existent image
        builder._add_image_to_slide(slide, "/nonexistent/path/to/image.jpg")

        # Should not raise exception, just log error
        # Slide should still be valid
        assert slide is not None

    def test_empty_presentation(self, builder, temp_output_dir):
        """Test building presentation with only title slide."""
        metadata = PresentationMetadata(
            topic="最小演示",
            audience="观众",
            purpose="目的"
        )

        slides = [
            Slide(
                slide_number=1,
                title="仅标题",
                slide_type="title",
                bullet_points=[BulletPoint(text="副标题", level=1)]
            )
        ]

        outline = PresentationOutline(metadata=metadata, slides=slides)
        output_path = str(Path(temp_output_dir) / "minimal.pptx")
        builder.build(outline, output_path)

        prs = Presentation(output_path)
        assert len(prs.slides) == 1

    def test_large_presentation(self, builder, temp_output_dir):
        """Test building a large presentation."""
        metadata = PresentationMetadata(
            topic="大型演示",
            audience="观众",
            purpose="目的"
        )

        slides = [
            Slide(
                slide_number=1,
                title="标题",
                slide_type="title",
                bullet_points=[BulletPoint(text="副标题", level=1)]
            )
        ]

        # Add many content slides
        for i in range(2, 31):
            slides.append(
                Slide(
                    slide_number=i,
                    title=f"内容页 {i}",
                    slide_type="content",
                    bullet_points=[
                        BulletPoint(text=f"要点 {j}", level=1)
                        for j in range(1, 6)
                    ]
                )
            )

        outline = PresentationOutline(metadata=metadata, slides=slides)
        output_path = str(Path(temp_output_dir) / "large.pptx")
        builder.build(outline, output_path)

        prs = Presentation(output_path)
        assert len(prs.slides) == 30
