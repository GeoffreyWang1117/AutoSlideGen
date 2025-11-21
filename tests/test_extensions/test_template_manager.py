"""
Tests for template manager extension.
"""

import pytest
import tempfile
import shutil
from pathlib import Path
from unittest.mock import Mock, patch
from pptx import Presentation

from autoslidegen.extensions.template_manager import (
    TemplateManager,
    CustomizableTemplateBuilder
)


class TestTemplateManager:
    """Tests for TemplateManager."""

    @pytest.fixture
    def temp_dir(self):
        """Create temporary directory."""
        temp_dir = tempfile.mkdtemp()
        yield temp_dir
        shutil.rmtree(temp_dir)

    @pytest.fixture
    def template_manager(self, temp_dir):
        """Create TemplateManager instance."""
        return TemplateManager(templates_dir=temp_dir)

    @pytest.fixture
    def sample_template(self, temp_dir):
        """Create a sample PPTX template."""
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        template_path = Path(temp_dir) / 'sample_template.pptx'
        prs.save(str(template_path))
        return str(template_path)

    def test_initialization(self, template_manager):
        """Test TemplateManager initialization."""
        assert Path(template_manager.templates_dir).exists()
        assert isinstance(template_manager.templates, dict)

    def test_add_template(self, template_manager, sample_template):
        """Test adding a template."""
        template_manager.add_template(sample_template, 'test_template')

        assert 'test_template' in template_manager.templates
        assert Path(template_manager.templates['test_template']).exists()

    def test_add_template_auto_name(self, template_manager, sample_template):
        """Test adding template with auto-generated name."""
        template_manager.add_template(sample_template)

        # Should use filename as template name
        assert 'sample_template' in template_manager.templates

    def test_remove_template(self, template_manager, sample_template):
        """Test removing a template."""
        template_manager.add_template(sample_template, 'test_template')
        result = template_manager.remove_template('test_template')

        assert result is True
        assert 'test_template' not in template_manager.templates

    def test_remove_nonexistent_template(self, template_manager):
        """Test removing a template that doesn't exist."""
        result = template_manager.remove_template('nonexistent')
        assert result is False

    def test_list_templates(self, template_manager, sample_template):
        """Test listing templates."""
        template_manager.add_template(sample_template, 'template1')

        templates = template_manager.list_templates()
        assert isinstance(templates, list)
        assert 'template1' in templates

    def test_get_template_path(self, template_manager, sample_template):
        """Test getting template path."""
        template_manager.add_template(sample_template, 'test_template')

        path = template_manager.get_template_path('test_template')
        assert path is not None
        assert Path(path).exists()

    def test_get_nonexistent_template_path(self, template_manager):
        """Test getting path for nonexistent template."""
        path = template_manager.get_template_path('nonexistent')
        assert path is None

    def test_get_template_info(self, template_manager, sample_template):
        """Test getting template information."""
        template_manager.add_template(sample_template, 'test_template')

        info = template_manager.get_template_info('test_template')
        assert info is not None
        assert 'name' in info
        assert 'path' in info
        assert 'num_layouts' in info

    def test_template_persistence(self, temp_dir, sample_template):
        """Test that templates persist across instances."""
        tm1 = TemplateManager(templates_dir=temp_dir)
        tm1.add_template(sample_template, 'persistent')

        # Create new instance with same directory
        tm2 = TemplateManager(templates_dir=temp_dir)
        assert 'persistent' in tm2.list_templates()


class TestCustomizableTemplateBuilder:
    """Tests for CustomizableTemplateBuilder."""

    @pytest.fixture
    def temp_dir(self):
        """Create temporary directory."""
        temp_dir = tempfile.mkdtemp()
        yield temp_dir
        shutil.rmtree(temp_dir)

    @pytest.fixture
    def template_manager(self, temp_dir):
        """Create TemplateManager with sample template."""
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        template_path = Path(temp_dir) / 'template.pptx'
        prs.save(str(template_path))

        tm = TemplateManager(templates_dir=temp_dir)
        tm.add_template(str(template_path), 'test_template')
        return tm

    @pytest.fixture
    def template_builder(self, template_manager):
        """Create CustomizableTemplateBuilder instance."""
        return CustomizableTemplateBuilder(template_manager, 'test_template')

    def test_initialization(self, template_builder):
        """Test CustomizableTemplateBuilder initialization."""
        assert template_builder.template_name == 'test_template'
        assert template_builder.template_manager is not None

    def test_set_layout_mapping(self, template_builder):
        """Test setting custom layout mapping."""
        mapping = {'title': 0, 'content': 1, 'section': 2}
        template_builder.set_layout_mapping(mapping)

        assert template_builder.layout_mapping == mapping

    def test_create_presentation(self, template_builder):
        """Test creating presentation from template."""
        prs = template_builder.create_presentation()

        assert isinstance(prs, Presentation)
        assert len(prs.slide_layouts) > 0

    def test_create_with_nonexistent_template(self, template_manager):
        """Test creating builder with nonexistent template."""
        with pytest.raises(Exception):
            CustomizableTemplateBuilder(template_manager, 'nonexistent')

    def test_get_layout_by_type(self, template_builder):
        """Test getting layout by type."""
        template_builder.set_layout_mapping({'title': 0})
        prs = template_builder.create_presentation()

        layout = template_builder.get_layout(prs, 'title')
        assert layout is not None

    def test_default_layout_mapping(self, template_builder):
        """Test default layout mapping."""
        assert 'title' in template_builder.layout_mapping
        assert 'content' in template_builder.layout_mapping
